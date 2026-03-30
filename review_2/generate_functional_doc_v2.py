import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_bold_run(paragraph, text):
    r = paragraph.add_run(text)
    r.bold = True
    return r

def create_functional_document():
    output_dir = r"E:\Projects\collabcodehub\review_2\review_2_docs"
    output_path = os.path.join(output_dir, "review_2_functional_document_v2.docx")
    
    # Delete the previously made one if it exists
    if os.path.exists(output_path):
        try:
            os.remove(output_path)
            print(f"Deleted old document at {output_path}")
        except Exception as e:
            print(f"Failed to delete old document: {e}")

    doc = Document()

    # Base font setup for defaults
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # Title
    doc.add_paragraph('\n')
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('Functional Document')
    run.bold = True
    run.font.size = Pt(14)
    doc.add_paragraph('\n')
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_run = subtitle.add_run('CollabCodeHub - Real-Time Collaborative Cloud Coding Workspace')
    s_run.bold = True
    s_run.font.size = Pt(12)
    doc.add_paragraph('\n')

    # 1. Introduction
    p = doc.add_paragraph()
    p.add_run('1. Introduction').bold = True
    doc.add_paragraph("The CollabCodeHub project aims to revolutionize remote software engineering by integrating real-time collaboration, autonomous AI assistance, and agile task tracking into a unified web-native workspace. Sprint 1 is focused on implementing the core Conflict-free Replicated Data Type (CRDT) engine for ultra-low latency code synchronization and securing workspace boundaries via robust Row-Level Security (RLS).")

    # 2. Product Goal
    p = doc.add_paragraph()
    p.add_run('2. Product Goal').bold = True
    doc.add_paragraph("The primary goal of this sprint is to deliver a highly robust, synchronized code editor that gracefully handles concurrent multi-user edits without central merge conflicts, alongside an isolated and secure room-based environment. This contributes directly to the overarching objective of providing a seamless, highly productive remote pair-programming experience.")

    # 3. Demography
    p = doc.add_paragraph()
    p.add_run('3. Demography (Users, Location)').bold = True
    
    doc.add_paragraph("Users")
    doc.add_paragraph("Target Users: Software Developers, Computer Science Students, Technical Educators, and Agile Project Teams.")
    doc.add_paragraph("User Characteristics: Ranging from beginner programming students needing remote lab assistance to senior engineers executing technical interviews or pair programming. Users possess varying levels of technical proficiency but require seamless remote access.")
    
    doc.add_paragraph("Location")
    doc.add_paragraph("Target Location: Worldwide / Global usage, highly optimized for geographically distributed remote teams operating across independent network connections.")

    # 4. Business Processes
    p = doc.add_paragraph()
    p.add_run('4. Business Processes').bold = True
    doc.add_paragraph("The key business processes orchestrated within the platform include:")
    
    doc.add_paragraph("Workspace Initialization and Authentication:")
    doc.add_paragraph("Process for users to securely log in via Supabase Auth, manage their profiles, and provision isolated collaborative coding rooms.")
    
    doc.add_paragraph("Real-Time Code Synchronization:")
    doc.add_paragraph("Process of broadcasting syntax-highlighted code mutations across decentralized peer nodes using Yjs CRDTs without operational transformation lag.")
    
    doc.add_paragraph("Agile Task Administration:")
    doc.add_paragraph("Process for users to define, assign, and track Kanban-style task cards synchronously within the coding environment without external context switching.")
    
    doc.add_paragraph("Contextual AI Code Generation:")
    doc.add_paragraph("Process for developers to query an embedded LLM (Qwen2.5-7B) to autonomously generate and apply code directly to the shared file tree.")

    # 5. Features
    p = doc.add_paragraph()
    p.add_run('5. Features').bold = True
    doc.add_paragraph("This sprint will focus on implementing the following key features:")
    
    doc.add_paragraph("Feature #1: Real-Time Concurrent Code Editor").runs[0].bold = True
    p1 = doc.add_paragraph()
    p1.add_run("1. Description: ").bold = True
    p1.add_run("Integration of the Monaco Editor tied dynamically to the Yjs shared object model. This allows multiple remote cursors to edit the same file natively without operational latency or file locks.")
    p2 = doc.add_paragraph()
    p2.add_run("2. User Story: ").bold = True
    p2.add_run("As a remote developer, I want to edit code simultaneously with my teammate so that we can conduct an effective peer-programming session without merge conflicts or screen-sharing lag.")
    
    doc.add_paragraph("Feature #2: Secure Room Management and Authorization").runs[0].bold = True
    p3 = doc.add_paragraph()
    p3.add_run("1. Description: ").bold = True
    p3.add_run("A fully protected workspace layer utilizing Supabase PostgreSQL Row-Level Security (RLS) policies. Only cryptographically verified room members can read/write data, tasks, or code within their specific room ID.")
    p4 = doc.add_paragraph()
    p4.add_run("2. User Story: ").bold = True
    p4.add_run("As an educator hosting a private coding interview, I want to restrict workspace access exclusively to invited candidates so that proprietary test questions and code remain confidential.")

    doc.add_paragraph("Feature #3: Integrated Agile Task Board").runs[0].bold = True
    p5 = doc.add_paragraph()
    p5.add_run("1. Description: ").bold = True
    p5.add_run("A live Kanban board integrated directly into the workspace UI. Task states (To-Do, In-Progress, Done) are synchronized immediately across all connected clients via Supabase Realtime channels.")
    p6 = doc.add_paragraph()
    p6.add_run("2. User Story: ").bold = True
    p6.add_run("As a project team lead, I want to create and assign tasks within the coding interface so that my developers can track their sprint progress without switching browser tabs.")

    # 6. Authorization Matrix
    p = doc.add_paragraph()
    p.add_run('6. Authorization Matrix').bold = True
    doc.add_paragraph("Define the roles and their corresponding access levels:")
    
    # Table creation
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Role'
    hdr_cells[1].text = 'Access Level'
    hdr_cells[0].paragraphs[0].runs[0].bold = True
    hdr_cells[1].paragraphs[0].runs[0].bold = True
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'Room Admin / Creator'
    row_cells[1].text = 'Full access to workspace settings, member invitation, user management, and all read/write editing privileges'
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'Workspace Member'
    row_cells[1].text = 'Access to real-time code execution, AI Agent invocation, task creation, and chat participation'
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'Authenticated User'
    row_cells[1].text = 'Limited access to their personal dashboard to view owned rooms or join new rooms via valid slug'
    
    row_cells = table.add_row().cells
    row_cells[0].text = 'Unauthenticated Guest'
    row_cells[1].text = 'No access to platform resources; forcefully redirected to the login/registration gateway'
    
    doc.add_paragraph('\n')

    # 7. Assumptions
    p = doc.add_paragraph()
    p.add_run('7. Assumptions').bold = True
    doc.add_paragraph("The development environment and infrastructure will remain stable throughout the sprint.")
    doc.add_paragraph("The Supabase Cloud and Liveblocks WebSockets infrastructure will maintain stable uptime, providing underlying connectivity.")
    doc.add_paragraph("Users possess a modern web browser capable of supporting WebSockets and Web Workers (specifically for Monaco compilation).")
    doc.add_paragraph("Stakeholders, including educators and project leads, will be available for feedback and clarification regarding room policies.")
    doc.add_paragraph("Team members possess the necessary knowledge of TypeScript, React hooks, and CRDT synchronization paradigms to complete the assigned tasks.")

    os.makedirs(output_dir, exist_ok=True)
    doc.save(output_path)
    print(f"Document successfully created and saved at {output_path}")

if __name__ == '__main__':
    create_functional_document()
