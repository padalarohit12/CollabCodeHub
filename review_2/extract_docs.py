import os
import glob
from docx import Document
import openpyxl

def extract_docx(file_path):
    print(f"Extracting docx: {file_path}")
    doc = Document(file_path)
    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text)
    
    # Also extract tables
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.strip().replace('\n', ' '))
            text.append(" | ".join(row_data))
    return "\n".join(text)

def extract_xlsx(file_path):
    print(f"Extracting xlsx: {file_path}")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text = []
    for sheet in wb.worksheets:
        text.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            row_data = [str(cell) if cell is not None else "" for cell in row]
            if any(row_data):
                text.append(" | ".join(row_data))
    return "\n".join(text)

def main():
    directory = r"E:\Projects\collabcodehub\review_2"
    output_file = os.path.join(directory, "extracted_texts.txt")
    
    all_text = []
    for filepath in glob.glob(os.path.join(directory, "*.docx")):
        try:
            content = extract_docx(filepath)
            all_text.append(f"================\nFILE: {os.path.basename(filepath)}\n================\n{content}\n\n")
        except Exception as e:
            all_text.append(f"Error extracting {filepath}: {e}\n")
            
    for filepath in glob.glob(os.path.join(directory, "*.xlsx")):
        try:
            content = extract_xlsx(filepath)
            all_text.append(f"================\nFILE: {os.path.basename(filepath)}\n================\n{content}\n\n")
        except Exception as e:
            all_text.append(f"Error extracting {filepath}: {e}\n")
            
    with open(output_file, "w", encoding="utf-8") as f:
        f.write("\n".join(all_text))
        
    print(f"Extraction complete. Output saved to {output_file}")

if __name__ == "__main__":
    main()
