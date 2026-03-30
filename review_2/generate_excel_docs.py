import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def create_functional_test_cases(output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Functional Test Cases"

    # Define Styles
    green_fill = PatternFill(start_color="549E83", end_color="549E83", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True, name="Times New Roman", size=11)
    regular_font = Font(name="Times New Roman", size=10)
    center_aligned = Alignment(horizontal="center", vertical="center", wrap_text=True)
    top_aligned = Alignment(horizontal="left", vertical="top", wrap_text=True)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Title Row
    ws.merge_cells('A1:G1')
    title_cell = ws['A1']
    title_cell.value = "Functional Test Case Template"
    title_cell.fill = green_fill
    title_cell.font = white_font
    title_cell.alignment = center_aligned
    title_cell.border = border

    # Headers
    headers = ["Feature", "Test Case", "Steps to execute test case", "Expected Output", "Actual Output", "Status", "More Information"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.value = header
        cell.fill = green_fill
        cell.font = white_font
        cell.alignment = center_aligned
        cell.border = border

    # Set Column Widths
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 35
    ws.column_dimensions['E'].width = 35
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 35

    # Data Rows
    data = [
        [
            "User Authentication", "Valid User Login", 
            "Open the application's login page.\n\nEnter a valid email.\n\nEnter a valid password.\n\nClick on the 'Login' button.", 
            "The user should be successfully logged into the system.\n\nThe application should redirect the user to the workspace dashboard.",
            "The user is successfully logged in.\n\nThe application redirects the user to the workspace dashboard.",
            "Pass",
            "No error messages are displayed.\n\nThe user profile information is correctly synced from Supabase."
        ],
        [
            "Real-time Code Collaboration", "Concurrent Multi-User Editing", 
            "User A and User B join the same room.\n\nUser A types code snippet A.\n\nUser B types code snippet B simultaneously on different lines.", 
            "Both edits should merge instantly without conflicts.\n\nBoth users must see the exact same code state instantly.",
            "Edits merge successfully with <50ms latency.\n\nYjs CRDT engine handles the conflict resolution perfectly.",
            "Pass",
            "Liveblocks WebSocket connection remains stable.\n\nNo operational transformation locks observed."
        ],
        [
            "AI Agent Integration", "Autonomous Code Generation",
            "Open AI pane in workspace.\n\nType 'Create a React Button component'.\n\nClick Send to HuggingFace API.",
            "AI processes prompt and returns JSON.\n\nThe system automatically inserts the generated code into a new file in the file tree.",
            "Code is generated and inserted correctly into the directory.\n\nChanges are broadcasted to all peers automatically.",
            "Pass",
            "Verify that Qwen2.5 model follows strict JSON formatting.\n\nVerify that aiAgent.ts applies modifications to Yjs local document."
        ],
        [
            "Room Security (RLS)", "Unauthorized Room Access",
            "Login as User C (Not invited to Room 1).\n\nAttempt to navigate to Room 1 URL directly.\n\nAttempt to fetch tasks for Room 1 via API console.",
            "The application should block access to the room UI.\n\nThe Supabase database should reject the API query with an empty array.",
            "Access to UI is blocked and redirected to dashboard.\n\nSupabase RLS strictly returns [] for unauthorized task fetches.",
            "Pass",
            "Ensure the PostgreSQL helper function 'is_room_member(r_id)' is successfully intercepting queries."
        ]
    ]

    for row_idx, row_data in enumerate(data, 3):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = value
            cell.font = regular_font
            cell.alignment = top_aligned
            cell.border = border

    wb.save(output_path)
    print(f"Saved Test Cases to {output_path}")

def create_retrospective(output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sprint Retrospective"

    # Define Styles
    blue_fill_title = PatternFill(start_color="00A2E8", end_color="00A2E8", fill_type="solid")
    blue_fill_header = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")
    light_blue_fill = PatternFill(start_color="99D9EA", end_color="99D9EA", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True, name="Times New Roman", size=11)
    black_font_bold = Font(color="000000", bold=True, name="Times New Roman", size=11)
    black_font_italic = Font(color="000000", italic=True, name="Times New Roman", size=10)
    regular_font = Font(name="Times New Roman", size=10)
    center_aligned = Alignment(horizontal="center", vertical="center", wrap_text=True)
    top_aligned = Alignment(horizontal="left", vertical="top", wrap_text=True)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Title Row
    ws.merge_cells('A1:D1')
    title_cell = ws['A1']
    title_cell.value = "Sprint Retrospective"
    title_cell.fill = blue_fill_title
    title_cell.font = black_font_bold
    title_cell.alignment = center_aligned
    title_cell.border = border

    # Guidelines Cell (E1:E2 merged)
    ws.merge_cells('E1:E2')
    g_cell = ws['E1']
    g_cell.value = "Guidelines"
    g_cell.fill = light_blue_fill
    g_cell.font = black_font_bold
    g_cell.alignment = center_aligned
    g_cell.border = border

    # Headers
    headers = ["What went well", "What went poorly", "What ideas do you have", "How should we take action"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.value = header
        cell.fill = blue_fill_header
        cell.font = black_font_bold
        cell.alignment = center_aligned
        cell.border = border

    # Guidelines Text Row
    guidelines = [
        "This section highlights the successes and positive outcomes from the sprint. It helps the team recognize achievements and identify practices that should be continued.",
        "This section identifies the challenges, roadblocks, or failures encountered during the sprint. It helps pinpoint areas that need improvement or change.",
        "This section is for brainstorming new approaches, tools, or strategies to enhance the team's efficiency, productivity, or project outcomes.",
        "This section outlines specific steps or solutions to address the issues and implement the ideas discussed, ensuring continuous improvement in future sprints."
    ]
    for col_num, text in enumerate(guidelines, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = text
        cell.fill = light_blue_fill
        cell.font = black_font_italic
        cell.alignment = top_aligned
        cell.border = border

    ws.cell(row=3, column=5).value = ""
    ws.cell(row=3, column=5).border = border

    # Set Column Widths
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 35
    ws.column_dimensions['E'].width = 15

    # Data Rows (Example/Actual Data)
    data = [
        [
            "Example : All tasks were completed on time.\nTeam communication was seamless.",
            "Requirements changed mid-sprint.",
            "Plan for a buffer to handle scope changes.",
            "Set stricter deadlines for finalizing requirements.",
            "Example"
        ],
        [
            "Real-time CRDT sync via Yjs integrated flawlessly with the Monaco Editor, providing <50ms latency across test regions.\n\nSupabase Row-Level Security effectively blocked unauthorized room access in all test scenarios, validating the zero-trust architecture.",
            "AI Agent (Qwen2.5) occasionally returned raw markdown blocks instead of pure JSON, causing the 'aiAgent.ts' parser to fail.\n\nMemory consumption spiked slowly when multiple large files were kept open simultaneously due to Yjs CRDT history bloat.",
            "Implement a more aggressive regex or a structured output parser for the HuggingFace API response to strip markdown wrappers natively.\n\nAdd automated state-vector garbage collection for Yjs documents to prevent memory leaks during extended sessions.",
            "Update 'aiAgent.ts' with a regex-based fallback parser to extract JSON blocks reliably before the next sprint review.\n\nResearch Yjs state-vector optimization and implement a periodic commit to Supabase to clear memory.",
            ""
        ]
    ]

    for idx, row_data in enumerate(data, 4):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=idx, column=col_idx)
            cell.value = value
            cell.font = regular_font
            cell.alignment = top_aligned
            cell.border = border

    wb.save(output_path)
    print(f"Saved Retrospective to {output_path}")

if __name__ == '__main__':
    out_dir = r"E:\Projects\collabcodehub\review_2\review_2_docs"
    os.makedirs(out_dir, exist_ok=True)
    tc_path = os.path.join(out_dir, "review_2_functional_test_cases.xlsx")
    retro_path = os.path.join(out_dir, "review_2_sprint_retrospective.xlsx")
    
    create_functional_test_cases(tc_path)
    create_retrospective(retro_path)
